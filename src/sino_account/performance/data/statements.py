"""Parse broker statement (對帳單) and inventory (庫存) Excel exports.

These two files supply the missing buy/sell history and the authoritative
current-holdings snapshot that the Shioaji API does not expose for past dates:

- 對帳單.xlsx: every executed trade (成交日, 商品, 買賣, 數量, 金額 ...).
- 庫存.xlsx:   current holdings snapshot (商品, 今日餘額 = shares, 現值 ...).

With both, daily holdings for any date inside the statement window can be
reconstructed exactly by walking the current inventory backward through trades.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

STATEMENT_FILENAME = "對帳單.xlsx"
INVENTORY_FILENAME = "庫存.xlsx"

# 買賣 -> share-count sign. Buys add shares (long) / cover shorts; sells remove.
_TRADE_SIGN = {
    "現買": 1,
    "券買": 1,
    "現賣": -1,
    "券賣": -1,
}


def _project_root() -> Path:
    # statements.py -> data -> performance -> sino_account -> src -> <root>
    return Path(__file__).resolve().parents[4]


def _resolve_path(path: str | Path | None, filename: str) -> Path:
    if path is not None:
        candidate = Path(path)
        if candidate.is_dir():
            candidate = candidate / filename
        if not candidate.exists():
            raise FileNotFoundError(f"找不到檔案: {candidate}")
        return candidate

    root = _project_root()
    for base in (root, root / "data"):
        candidate = base / filename
        if candidate.exists():
            return candidate
    raise FileNotFoundError(
        f"找不到 {filename}（已搜尋專案根目錄與 data/）。請將檔案放在專案根目錄。"
    )


def _normalize_date(value: Any) -> str:
    """Normalize 成交日 (e.g. '2025/07/01' or a datetime) to 'YYYY-MM-DD'."""
    if value is None:
        return ""
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip().replace("/", "-")
    return text[:10]


def _split_code_name(value: Any) -> tuple[str, str]:
    """'00830 國泰費城半導體' -> ('00830', '國泰費城半導體')."""
    text = str(value or "").strip()
    if not text:
        return "", ""
    parts = text.split(None, 1)
    code = parts[0]
    name = parts[1].strip() if len(parts) > 1 else ""
    return code, name


def _header_index(rows: list[tuple[Any, ...]]) -> dict[str, int]:
    header = rows[0] if rows else ()
    return {str(cell).strip(): idx for idx, cell in enumerate(header) if cell is not None}


def _to_float(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def load_statement(path: str | Path | None = None) -> list[dict[str, Any]]:
    """Return executed trades as normalized rows.

    Each row: {date, code, name, action, shares, price, amount, cashflow, delta}.
    - delta:    signed share change (+buy / -sell), used for backward holdings replay.
    - cashflow: signed cash change (+應收 on sell, -應付 on buy), trade-driven only.
    """
    statement_path = _resolve_path(path, STATEMENT_FILENAME)
    workbook = load_workbook(statement_path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()

    if not rows:
        return []

    col = _header_index(rows)
    date_i = col.get("成交日", 0)
    product_i = col.get("商品", 1)
    action_i = col.get("買賣", 2)
    qty_i = col.get("數量", 3)
    price_i = col.get("成交價", 4)
    amount_i = col.get("價金", 5)
    payable_i = col.get("應付金額", 8)
    receivable_i = col.get("應收金額", 9)

    trades: list[dict[str, Any]] = []
    for raw in rows[1:]:
        if not raw:
            continue
        first = str(raw[date_i] if date_i < len(raw) else "").strip()
        if not first or first == "合計":
            continue

        date = _normalize_date(raw[date_i] if date_i < len(raw) else "")
        code, name = _split_code_name(raw[product_i] if product_i < len(raw) else "")
        action = str(raw[action_i] if action_i < len(raw) else "").strip()
        if not code or not date:
            continue

        sign = _TRADE_SIGN.get(action)
        if sign is None:
            continue

        shares = _to_float(raw[qty_i] if qty_i < len(raw) else 0)
        payable = _to_float(raw[payable_i] if payable_i < len(raw) else 0)
        receivable = _to_float(raw[receivable_i] if receivable_i < len(raw) else 0)

        trades.append(
            {
                "date": date,
                "code": code,
                "name": name,
                "action": action,
                "shares": shares,
                "price": _to_float(raw[price_i] if price_i < len(raw) else 0),
                "amount": _to_float(raw[amount_i] if amount_i < len(raw) else 0),
                "delta": sign * shares,
                "cashflow": receivable - payable,
            }
        )

    trades.sort(key=lambda row: row["date"])
    return trades


def load_inventory(path: str | Path | None = None) -> dict[str, Any]:
    """Return current holdings snapshot keyed by code.

    {
      "holdings": {code: shares},   # 今日餘額
      "names": {code: name},
      "market_value": {code: 現值},
      "total_market_value": float,
    }
    """
    inventory_path = _resolve_path(path, INVENTORY_FILENAME)
    workbook = load_workbook(inventory_path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()

    if not rows:
        return {"holdings": {}, "names": {}, "market_value": {}, "total_market_value": 0.0}

    col = _header_index(rows)
    product_i = col.get("商品", 0)
    today_i = col.get("今日餘額", 6)
    value_i = col.get("現值", 9)

    holdings: dict[str, float] = {}
    names: dict[str, str] = {}
    market_value: dict[str, float] = {}

    for raw in rows[1:]:
        if not raw:
            continue
        product = str(raw[product_i] if product_i < len(raw) else "").strip()
        if not product or product == "合計":
            continue

        code, name = _split_code_name(product)
        if not code:
            continue

        shares = _to_float(raw[today_i] if today_i < len(raw) else 0)
        value = _to_float(raw[value_i] if value_i < len(raw) else 0)
        holdings[code] = holdings.get(code, 0.0) + shares
        names[code] = name
        market_value[code] = market_value.get(code, 0.0) + value

    return {
        "holdings": holdings,
        "names": names,
        "market_value": market_value,
        "total_market_value": sum(market_value.values()),
    }


def reconstruct_holdings_on(
    inventory_now: dict[str, float],
    trades: list[dict[str, Any]],
    as_of: str,
) -> dict[str, float]:
    """Holdings (shares per code) at the close of `as_of`.

    holdings(D) = inventory_now - sum(delta for trades executed after D).
    Trades on `as_of` itself are already reflected in that day's close.
    """
    holdings: dict[str, float] = {code: float(s) for code, s in inventory_now.items()}
    for trade in trades:
        if trade["date"] > as_of:
            code = trade["code"]
            holdings[code] = holdings.get(code, 0.0) - trade["delta"]
    return {code: shares for code, shares in holdings.items() if abs(shares) > 1e-9}


def reconstruct_cash_on(
    cash_now: float,
    trades: list[dict[str, Any]],
    as_of: str,
) -> float:
    """Trade-only cash at the close of `as_of` (ignores dividends/deposits/withdrawals)."""
    cash = cash_now
    for trade in trades:
        if trade["date"] > as_of:
            cash -= trade["cashflow"]
    return cash
